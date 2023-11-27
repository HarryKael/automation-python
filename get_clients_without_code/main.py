import pyodbc as pd
import pandas
from pandas import DataFrame
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook.workbook import Workbook
from openpyxl.cell.cell import Cell
from pyodbc import Connection, Cursor
from password_d import password
from test.test_changes import create_worksheet, getDocument

def make_connection() -> Connection:
    DRIVER_NAME = 'SQL SERVER'
    SERVER_NAME = "SERVER'S NAME"
    DATABASE_NAME = "DATABASE' NAME"
    
    connection_string = """
        DRIVER={{{DRIVER_NAME}}};
        SERVER={{{SERVER_NAME}}};
        DATABASE={{{DATABASE_NAME}}};
        Trust_Connection=yes;
        uid=sist-hisrael;
        pwd={{{password}}};
    """ 
    cnxn : Connection = pd.connect(
        driver = DRIVER_NAME,
        server = SERVER_NAME,
        database = DATABASE_NAME,
        trust_connection = 'yes',
        # uid = ,
        pwd = password,
    )
    return cnxn

def get_data(row_value, cnxn) -> DataFrame:
    query = f"""SELECT m.CODCLI
                FROM MAECLIENTE m
                WHERE m.RIFCLI = '{row_value}'"""
    return pandas.read_sql(query, cnxn)

def get_municipio(municipio, cnxn) -> DataFrame:
    query = f"""SELECT c.CODESTADO
                FROM CTRMUNICIPIOS c
                WHERE c.MUNICIPIO = '{municipio}';"""
    return pandas.read_sql(query, cnxn)

def make_query(row_value: str, cnxn, function):
    # ! Do the query.
    data: DataFrame = function(row_value, cnxn)
    return data.values

def main(working_book):
    sheet2: str = 'Sin codcli'
    necessary_fields = ['Nombre', 'RNC', 'Dirección', 'Provincia', 'Municipio', 'Sector', 'CODESTADO', 'CODCLI', 'EMAIL', 'PROPIETA', 'CIPROPIE', 'TLFPROPIE', 'DIRPROPIE', 'REGENTE', 'CIREG', 'TLFREG', 'DIRREG', 'CIUDAD', 'CODIGOPOSTAL',]
    cnxn = make_connection()
    workbook: Workbook = openpyxl.load_workbook(getDocument())
    i = 0
    municipios: list[tuple[str]] = []
    
    worksheet: Worksheet = workbook.get_sheet_by_name(working_book)
    worksheet2: Worksheet = create_worksheet(workbook, sheet2)
    
    worksheet2.append(necessary_fields)
    for row in worksheet:
        i += 1
        row: Cell = row
        row_value: str = str(row[3].value).strip()
        municipio = str(row[6].value).strip()
        municipio: str = municipio.replace('CENTRO (DN)', '').strip()

        if row_value != 'None' and row_value != 'RNC':
            values2 = make_query(row_value, cnxn, get_data)
            if len(values2) == 0:
                values_codestado = make_query(municipio, cnxn, get_municipio)
                municipios.append((str(municipio), str(values_codestado)))
                datas = [str(row[1].value).strip(), row_value, str(row[4].value).strip(), str(row[5].value).strip(), str(row[6].value).strip(), str(row[7].value).strip(), str(values_codestado)]
                worksheet2.append(datas)

    workbook.save(getDocument('new'))
    for municipio in municipios:
        print(municipio)

if __name__ == '__main__':
    main('Listado clientes')
