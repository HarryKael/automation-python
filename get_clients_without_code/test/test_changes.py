import openpyxl
from openpyxl.cell import Cell
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
import unittest

def create_worksheet(workbook: Workbook, title: str):
    workbook.create_sheet(title)
    return workbook.get_sheet_by_name(title)

def getDocument(new_name: str = '', v:int = 1) -> str:
    doc = ''
    match v:
        case 1:
            doc = './clientes' + new_name + '.xlsx'
        case 2:
            doc = './test/clientes' + new_name + '.xlsx'
    return doc

class TestChanges(unittest.TestCase):
    new_sheet: str = 'New'

    def test_creation(self):
        index_sheet2 = 0
        doc = getDocument('', 2)
        workbook = openpyxl.load_workbook(doc)
        worksheet = create_worksheet(workbook, self.new_sheet)
        index_sheet = workbook.get_index(worksheet)
        try:
            index_sheet2 = workbook.get_index(Worksheet(workbook, 'Hola'))
        except ValueError: pass
        self.assertNotEqual(index_sheet, 0)
        self.assertEqual(index_sheet2, 0)

        # ! Guardar
        doc = getDocument('new', 2)
        workbook.save(doc)
        self.assertLess(4, len(workbook.worksheets))

if __name__ == '__main__':
    unittest.main()